"""
make_ppt.py — generate advisor weekly-report PPTX from results figures
Usage: python make_ppt.py [scenario] [output.pptx]
"""
import sys, os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN

scen = sys.argv[1] if len(sys.argv) > 1 else "near-term"
out = sys.argv[2] if len(sys.argv) > 2 else f"Biochar_Progress_{scen}.pptx"
figdir = os.path.join("results_v2", scen, "figures")
W, H = Inches(13.333), Inches(7.5)

def add_slide(prs, title, subtitle=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    box = slide.shapes.add_textbox(Inches(0.4), Inches(0.25), Inches(12.5), Inches(0.9))
    tf = box.text_frame
    tf.text = title
    tf.paragraphs[0].font.size = Pt(28)
    tf.paragraphs[0].font.bold = True
    if subtitle:
        sub = slide.shapes.add_textbox(Inches(0.4), Inches(1.0), Inches(12.5), Inches(0.6))
        stf = sub.text_frame
        stf.text = subtitle
        stf.paragraphs[0].font.size = Pt(14)
    return slide

def add_pic(slide, path, left, top, width):
    if os.path.exists(path):
        slide.shapes.add_picture(path, Inches(left), Inches(top), width=Inches(width))
    else:
        print("  MISSING:", path)

def bullets(slide, items, left=0.6, top=1.7, width=5.6, size=16):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(5.4))
    tf = box.text_frame
    tf.word_wrap = True
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = it
        p.font.size = Pt(size)
        p.space_after = Pt(8)
    return box

prs = Presentation()
prs.slide_width, prs.slide_height = W, H

# S1 cover
s = add_slide(prs, "Carbon Policy Design for a Wisconsin Biochar Supply Chain",
              "Market Clearing, Endogenous Carbon Prices, and Technology Choice  |  Progress report, September 2026")
bullets(s, ["72 counties, 13 feedstocks, T1 drying + T2 pyrolysis (300/500°C), 39 technologies",
            "Three policy paradigms: baseline-and-credit | cap-and-trade (endogenous price) | tiered tax + VM0044 eligibility",
            "Data: DOE Billion-Ton 2023 (near-term / mature-market) + USDA NASS + USFS FIA"])

# S2 research questions
s = add_slide(prs, "Research Questions")
bullets(s, [
    "Q1: How should carbon policy enter a supply chain optimization model so that it genuinely changes locations, technology, and output?",
    "Q2: What is the ENDOGENOUS carbon price of such a system, and how does it compare with RGGI / EU ETS / VM0044 markets?",
    "Q3: Which policy design performs best on cost, removal, and system expansion?",
])

# S3 framework
s = add_slide(prs, "Modeling Framework (Sampat 2019 market clearing)")
add_pic(s, os.path.join(figdir, "fig0_superstructure.png"), 1.2, 1.6, 10.8)
bullets(s, ["Duals of mass balances = spatial inherent values (equilibrium prices under coordination)"],
        top=6.7, width=11.0, size=14)

# S4 baseline results
s = add_slide(prs, f"Baseline Market Equilibrium ({scen})")
bullets(s, [
    "Equilibrium: 2.30 Mt biochar/yr — H ($800) and M ($500) fully served; L ($250) unserved (marginal cost > price)",
    "System surplus $708M/yr (near-term) | $736M/yr (mature-market)",
    "Net removal: -1.60 Mt CO2e/yr (process +979 kt, transport +24 kt, removal -2,598 kt)",
    "Corn stover 5.9 Mt wet fully utilized; facilities T1=78, T2-300=38",
], width=6.2)
add_pic(s, os.path.join(figdir, "fig1_facility_map.png"), 7.0, 1.5, 5.6)

# S5 inherent values
s = add_slide(prs, "Market Clearing: Inherent Values from LP Duals")
add_pic(s, os.path.join(figdir, "fig4_iv_biochar.png"), 0.4, 1.5, 5.9)
add_pic(s, os.path.join(figdir, "fig5_iv_biomass.png"), 6.9, 1.5, 5.9)
bullets(s, ["Biochar inherent value $305-345/t (near-term) vs $260-288/t (mature): abundant supply flattens location rents",
            "Wet biomass value ($0-77/t) far below farmgate: system is demand-constrained"],
        top=6.6, width=12.2, size=13)

# S6 policy A
s = add_slide(prs, "Paradigm A: Baseline-and-Credit")
add_pic(s, os.path.join(figdir, "fig6_policyA.png"), 0.4, 1.5, 12.4)
bullets(s, ["Credits scale with throughput → policy moves the solution",
            "p_c=$25/t activates the bulk segment; $200/t mobilizes ~100% of near-term biomass (surplus $3.36B/yr)"],
        top=5.9, width=12.2, size=14)

# S7 policy B
s = add_slide(prs, "Paradigm B: Cap-and-Trade — Endogenous Carbon Price")
add_pic(s, os.path.join(figdir, "fig7_policyB_MAC.png"), 0.4, 1.5, 12.4)
bullets(s, ["Vertex dual of the emission cap = endogenous carbon price; MAC ≈ $351-391/tCO2e first binding intervals, rising to $599 (above EU ETS & voluntary prices)",
            "Net emission caps NEVER bind (π=0): net-negative systems need explicit removal valuation"],
        top=5.9, width=12.2, size=14)

# S8 policy C
s = add_slide(prs, "Paradigm C: Tiered Tax + VM0044 Eligibility")
add_pic(s, os.path.join(figdir, "fig8_policyC_tech.png"), 0.4, 1.5, 12.4)
bullets(s, ["H/C ≤ 0.7 gate: ALL 300°C biochar is torrefaction-like (H/C 0.9-1.4) → zero credits; 500°C (H/C 0.40-0.52) always qualifies",
            "Crediting converts the system to 500°C as p_c rises: 2.25 Mt near-term / 4.25 Mt mature at $200/t — credits buy durability, not tonnes",
            "Stress test (500°C H/C +0.25, beyond published range): conversion falls back to 0.31/1.59 Mt — robust within the 0.40-0.55 literature band"],
        top=5.9, width=12.2, size=14)

# S9 comparison
s = add_slide(prs, "Policy Comparison")
add_pic(s, os.path.join(figdir, "fig9_policy_compare.png"), 0.4, 1.5, 12.4)

# S10 sensitivity
s = add_slide(prs, "Sensitivity: Demand-side Dominance")
add_pic(s, os.path.join(figdir, "fig10_tornado.png"), 0.4, 1.5, 7.4)
bullets(s, ["Demand price ±20% → surplus ±39% (largest lever)",
            "Demand capacity -20% → quantity -20%",
            "Farmgate ±20% → ±9%; transport ±1% (supply-side minor)"],
        left=8.4, top=1.8, width=4.4, size=15)

# S11 next steps
s = add_slide(prs, "Next Steps")
bullets(s, [
    "Multi-period / seasonal inventory extension (advisor direction)",
    "BT23 wastes layer characterized (2.27 Mt at <=$70/t, dairy manure 1.13 Mt) for co-location analysis",
    "Tighten MIP gaps on HPC; VM0044 durability factors aligned (0.65/0.80, v1.2 Table 3)",
    "Sensitivity on adoption rate & quality premium calibration",
], width=12.2)

prs.save(out)
print(f"saved {out} ({len(prs.slides.slides if hasattr(prs.slides,'slides') else prs.slides._sldIdLst)} slides)")
